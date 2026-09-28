"""Shared export rules: Vietnam dates, stable caller order, safe spreadsheet cells."""
import csv
import io
import json
import pandas as pd
from utils.datetime_utils import DATE_FIELDS, format_vietnam, vietnam_datetime, localize_json

POST_FIELDS = ['rank', 'post_id', 'group_id', 'post_url', 'author_name', 'content', 'created_time', 'likes', 'comments', 'shares', 'engagement_score', 'collected_at']
COMMENT_FIELDS = ['rank', 'comment_id', 'post_id', 'comment_url', 'author_name', 'content', 'created_time', 'reactions', 'replies', 'engagement_score', 'collected_at', 'score_status', 'parent_comment_id', 'depth']

def safe(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return value

def csv_bytes(rows, fields):
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow(fields)
    for row in rows:
        writer.writerow([format_vietnam(row.get(k)) if k in DATE_FIELDS else safe(row.get(k)) for k in fields])
    return stream.getvalue().encode('utf-8-sig')

def excel_bytes(sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows, fields in sheets:
        ws = wb.create_sheet(name)
        ws.append(fields)
        for row in rows:
            values = []
            for key in fields:
                value = row.get(key)
                if key in DATE_FIELDS:
                    date = vietnam_datetime(value)
                    value = None if pd.isna(date) else date.to_pydatetime()
                else:
                    value = safe(value)
                values.append(ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value)
            ws.append(values)
        for col, key in enumerate(fields, 1):
            ws.cell(1, col).font = Font(bold=True)
            ws.column_dimensions[ws.cell(1, col).column_letter].width = 60 if 'content' in key else 24
            if key in DATE_FIELDS:
                for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    row[0].number_format = 'dd/mm/yyyy hh:mm:ss'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def json_bytes(data):
    return json.dumps(localize_json(data), ensure_ascii=False, indent=2).encode('utf-8')